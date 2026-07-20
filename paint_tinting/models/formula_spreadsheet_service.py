# -*- coding: utf-8 -*-
"""XLSX template generation, parsing, validation, and staged draft creation."""

import io
import json
import logging

from odoo import _, api, models
from odoo.exceptions import ValidationError


_logger = logging.getLogger(__name__)

FORMULA_HEADERS = [
    'Reference', 'Fandeck', 'Colour Code', 'Base Category', 'Base UoM', 'Base Attribute',
    *['C%s' % number for number in range(1, 17)],
    'Change Reason', 'Notes', 'Total Shots', 'Total ML',
]


class TintingFormulaSpreadsheetService(models.AbstractModel):
    _name = 'tinting.formula.spreadsheet.service'
    _description = 'Tinting Formula Spreadsheet Service'

    @api.model
    def build_template(self):
        try:
            import xlsxwriter
        except ImportError as exc:
            raise ValidationError(_('Python package xlsxwriter is required to generate the template.')) from exc

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        title = workbook.add_format({
            'bold': True, 'font_size': 16, 'font_color': '#FFFFFF',
            'bg_color': '#156B63', 'align': 'center', 'valign': 'vcenter',
        })
        section = workbook.add_format({
            'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#23877D',
            'align': 'center', 'valign': 'vcenter', 'border': 1,
        })
        header = workbook.add_format({
            'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#3A4D5C',
            'align': 'center', 'valign': 'vcenter', 'border': 1, 'text_wrap': True,
        })
        input_text = workbook.add_format({'bg_color': '#F7FBFA', 'border': 1})
        input_number = workbook.add_format({'bg_color': '#F7FBFA', 'border': 1, 'num_format': '0.00'})
        formula_fmt = workbook.add_format({'bg_color': '#E8F2F0', 'border': 1, 'num_format': '0.00'})
        note = workbook.add_format({'font_color': '#34495E', 'text_wrap': True, 'valign': 'top'})
        warning = workbook.add_format({'font_color': '#9C5700', 'bg_color': '#FFEB9C', 'text_wrap': True})

        data_sheet = workbook.add_worksheet('Formula Data')
        instructions = workbook.add_worksheet('Instructions')
        lookups = workbook.add_worksheet('Lookups')
        data_sheet.hide_gridlines(2)
        instructions.hide_gridlines(2)
        lookups.hide_gridlines(2)

        data_sheet.merge_range(0, 0, 0, len(FORMULA_HEADERS) - 1, 'Paint Tinting Formula Draft Import', title)
        data_sheet.set_row(0, 28)
        data_sheet.merge_range(1, 0, 1, 5, 'Formula Identity (must remain stable across versions)', section)
        data_sheet.merge_range(1, 6, 1, 21, 'Colorant Shots (enter values greater than zero)', section)
        data_sheet.merge_range(1, 22, 1, 23, 'Audit Details', section)
        data_sheet.merge_range(1, 24, 1, 25, 'Calculated Preview', section)
        for column, label in enumerate(FORMULA_HEADERS):
            data_sheet.write(2, column, label, header)
        data_sheet.freeze_panes(3, 6)
        data_sheet.autofilter(2, 0, 502, len(FORMULA_HEADERS) - 1)
        widths = [14, 16, 14, 22, 12, 22] + [9] * 16 + [28, 28, 13, 13]
        for index, width in enumerate(widths):
            data_sheet.set_column(index, index, width)
        for row in range(3, 503):
            for column in range(0, 6):
                data_sheet.write_blank(row, column, None, input_text)
            for column in range(6, 22):
                data_sheet.write_blank(row, column, None, input_number)
                data_sheet.data_validation(row, column, row, column, {
                    'validate': 'decimal', 'criteria': '>=', 'value': 0,
                    'input_title': 'Colorant shots',
                    'input_message': 'Blank or zero is ignored. Use at most two decimals.',
                    'error_title': 'Invalid shots',
                    'error_message': 'Shots must be zero or greater.',
                })
            data_sheet.write_blank(row, 22, None, input_text)
            data_sheet.write_blank(row, 23, None, input_text)
            excel_row = row + 1
            data_sheet.write_formula(row, 24, '=IF(COUNTA(G%d:V%d)=0,"",SUM(G%d:V%d))' % (
                excel_row, excel_row, excel_row, excel_row,
            ), formula_fmt)
            data_sheet.write_formula(row, 25, '=IF(Y%d="","",Y%d*0.616)' % (excel_row, excel_row), formula_fmt)

        lookup_sets = self._template_lookups()
        for column, (label, values) in enumerate(lookup_sets.items()):
            lookups.write(0, column, label, header)
            lookups.set_column(column, column, 28)
            for row, value in enumerate(values, start=1):
                lookups.write(row, column, value, input_text)
        lookup_columns = {'Fandecks': 1, 'Colour Codes': 2, 'Base Categories': 3, 'Base UoMs': 4, 'Base Attributes': 5}
        for lookup_name, data_column in lookup_columns.items():
            values = lookup_sets.get(lookup_name, [])
            if values:
                lookup_index = list(lookup_sets).index(lookup_name)
                first = xlsxwriter.utility.xl_col_to_name(lookup_index)
                formula = "='Lookups'!$%s$2:$%s$%s" % (first, first, len(values) + 1)
                data_sheet.data_validation(3, data_column, 502, data_column, {
                    'validate': 'list', 'source': formula,
                })

        instructions.merge_range('A1:F1', 'How to Import Tinting Formula Drafts', title)
        instructions.set_row(0, 28)
        instruction_rows = [
            ('1', 'Use one row per complete formula recipe.'),
            ('2', 'Identity is Colour Code + Base Category + Base UoM + normalized Base Attribute.'),
            ('3', 'If the identity is new, import creates Version 1 Draft.'),
            ('4', 'If the identity exists and shots changed, import creates the next Draft version.'),
            ('5', 'An identical recipe is detected as a duplicate and skipped.'),
            ('6', 'Blank and zero colorants are ignored. Negative shots are rejected.'),
            ('7', 'Imported versions remain Draft until a Formula Manager approves them.'),
            ('8', 'Never rename columns or worksheets. Total columns are calculated and ignored during import.'),
        ]
        instructions.write('A3', 'Step', header)
        instructions.write('B3', 'Instruction', header)
        for row, values in enumerate(instruction_rows, start=3):
            instructions.write(row, 0, values[0], section)
            instructions.write(row, 1, values[1], note)
        instructions.set_column('A:A', 9)
        instructions.set_column('B:B', 95)
        instructions.write('A13', 'Important', header)
        instructions.merge_range('B13:F15',
                                 'Approval is deliberately separate from import. Review the preview and recipe differences before bulk approval.',
                                 warning)

        workbook.close()
        output.seek(0)
        payload = output.read()
        _logger.info(
            "FORMULA_TEMPLATE_GENERATED bytes=%s lookups=%s user_id=%s",
            len(payload), {key: len(value) for key, value in lookup_sets.items()}, self.env.user.id,
        )
        return payload

    @api.model
    def _template_lookups(self):
        attributes = self.env['product.template.attribute.value'].search([]).mapped('name')
        service = self.env['tinting.formula.recipe.service']
        return {
            'Fandecks': sorted(set(self.env['colour.fandeck'].search([]).mapped('name'))),
            'Colour Codes': sorted(set(self.env['colour.code'].search([]).mapped('code'))),
            'Base Categories': sorted(set(self.env['product.category'].search([]).mapped('name'))),
            'Base UoMs': sorted(set(self.env['uom.uom'].search([]).mapped('name'))),
            'Base Attributes': sorted(set(filter(None, (service.normalize_attribute(v) for v in attributes)))),
            'Colorant Codes': ['C%s' % number for number in range(1, 17)],
        }

    @api.model
    def parse_to_batch(self, payload, filename):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValidationError(_('Python package openpyxl is required to import XLSX files.')) from exc
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
        except Exception as exc:
            _logger.exception("FORMULA_IMPORT_WORKBOOK_OPEN_FAILED filename=%s bytes=%s", filename, len(payload))
            raise ValidationError(_('The uploaded XLSX file could not be read: %s', exc)) from exc
        if 'Formula Data' not in workbook.sheetnames:
            raise ValidationError(_('Worksheet "Formula Data" was not found. Please use the provided template.'))
        sheet = workbook['Formula Data']
        header_row = next(sheet.iter_rows(
            min_row=3, max_row=3, min_col=1, max_col=len(FORMULA_HEADERS), values_only=True,
        ))
        actual_headers = [str(value or '').strip() for value in header_row]
        if actual_headers != FORMULA_HEADERS:
            _logger.warning("FORMULA_IMPORT_HEADERS_INVALID expected=%s actual=%s", FORMULA_HEADERS, actual_headers)
            raise ValidationError(_('The spreadsheet columns do not match the current formula template.'))

        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': __import__('base64').b64encode(payload),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'res_model': 'tinting.formula.import.batch',
        })
        batch = self.env['tinting.formula.import.batch'].create({
            'name': _('Formula Import - %s', filename),
            'filename': filename,
            'source_attachment_id': attachment.id,
        })
        attachment.write({'res_id': batch.id})
        row_values = []
        seen_identities = set()
        for row_number, row_values_tuple in enumerate(sheet.iter_rows(
            min_row=4, min_col=1, max_col=len(FORMULA_HEADERS), values_only=True,
        ), start=4):
            values = list(row_values_tuple)
            if not any(value not in (None, '') for value in values[:24]):
                continue
            row_values.append(self._validate_row(batch, row_number, values, seen_identities))
        if row_values:
            self.env['tinting.formula.import.row'].create(row_values)
        batch.log_summary = _(
            'Rows read: %(rows)s; valid: %(valid)s; duplicates: %(duplicates)s; errors: %(errors)s',
            rows=batch.row_count, valid=batch.valid_count,
            duplicates=batch.duplicate_count, errors=batch.error_count,
        )
        _logger.info(
            "FORMULA_IMPORT_PREVIEW_CREATED batch_id=%s filename=%s rows=%s valid=%s duplicates=%s errors=%s",
            batch.id, filename, batch.row_count, batch.valid_count,
            batch.duplicate_count, batch.error_count,
        )
        return batch

    @api.model
    def _validate_row(self, batch, row_number, values, seen_identities):
        service = self.env['tinting.formula.recipe.service']
        identity_values = [str(value or '').strip() for value in values[:6]]
        reference, fandeck_name, colour_code, category_name, uom_name, attribute_name = identity_values
        recipe = {}
        errors = []
        for offset, code in enumerate(['C%s' % number for number in range(1, 17)], start=6):
            raw = values[offset]
            if raw in (None, ''):
                continue
            try:
                shots = service.normalize_shots(raw)
                if shots < 0:
                    errors.append(_('%s cannot be negative.', code))
                elif shots > 0:
                    recipe[code] = shots
            except (TypeError, ValueError):
                errors.append(_('%(code)s has invalid shots "%(value)s".', code=code, value=raw))
        change_reason = str(values[22] or '').strip()
        notes = str(values[23] or '').strip()
        if not colour_code:
            errors.append(_('Colour Code is required.'))
        if not category_name:
            errors.append(_('Base Category is required.'))
        if not uom_name:
            errors.append(_('Base UoM is required.'))
        if not recipe:
            errors.append(_('At least one colorant must have shots greater than zero.'))

        resolved = {}
        if not errors:
            try:
                resolved = self._resolve_identity(fandeck_name, colour_code, category_name, uom_name, attribute_name)
            except ValidationError as exc:
                errors.append(str(exc))

        state = 'error' if errors else 'valid'
        planned_action = 'none' if errors else 'new_formula'
        message = '\n'.join(errors)
        family = False
        if not errors:
            family = service.find_family(
                resolved['colour'], resolved['category'], resolved['uom'], resolved['attribute_name'],
            )
            identity_key = (
                resolved['colour'].id, resolved['category'].id, resolved['uom'].id,
                resolved['attribute_name'], service.signature(recipe),
            )
            if identity_key in seen_identities:
                state, planned_action = 'duplicate', 'skip'
                message = _('Duplicate row within this spreadsheet.')
            else:
                seen_identities.add(identity_key)
                if family:
                    existing = family.version_ids.filtered(
                        lambda version: version.recipe_signature == service.signature(recipe)
                    )[:1]
                    legacy_matches = (
                        not family.version_ids
                        and family.formula_line_ids
                        and service.signature(family.get_legacy_shots_dict()) == service.signature(recipe)
                    )
                    if existing:
                        state, planned_action = 'duplicate', 'skip'
                        message = _('Recipe already exists as %s.', existing.display_name)
                    elif legacy_matches:
                        state, planned_action = 'duplicate', 'skip'
                        message = _('Recipe matches the existing legacy formula snapshot.')
                    else:
                        planned_action = 'new_version'
                        message = _('Will create the next Draft version of %s.', family.display_name)
                else:
                    message = _('Will create a new formula family and Version 1 Draft.')
        _logger.debug(
            "FORMULA_IMPORT_ROW_VALIDATED batch_id=%s row=%s state=%s action=%s family_id=%s errors=%s recipe=%s",
            batch.id, row_number, state, planned_action, family.id if family else False, errors, recipe,
        )
        return {
            'batch_id': batch.id,
            'row_number': row_number,
            'external_reference': reference,
            'fandeck_name': fandeck_name,
            'colour_code': colour_code,
            'base_category_name': category_name,
            'base_uom_name': uom_name,
            'base_attribute_name': service.normalize_attribute(attribute_name),
            'recipe_json': json.dumps(service.normalize_recipe(recipe), sort_keys=True),
            'change_reason': change_reason,
            'notes': notes,
            'state': state,
            'planned_action': planned_action,
            'message': message,
            'formula_id': family.id if family else False,
        }

    @api.model
    def _resolve_identity(self, fandeck_name, colour_code, category_name, uom_name, attribute_name):
        service = self.env['tinting.formula.recipe.service']
        fandeck = self.env['colour.fandeck'].search([('name', '=ilike', fandeck_name)], limit=1) if fandeck_name else False
        if fandeck_name and not fandeck:
            raise ValidationError(_('Fandeck "%s" was not found.', fandeck_name))
        colour_domain = [('code', '=ilike', colour_code)]
        if fandeck:
            colour_domain.append(('fandeck_id', '=', fandeck.id))
        colour_candidates = self.env['colour.code'].search(colour_domain)
        if not colour_candidates:
            raise ValidationError(_('Unknown colour code %(code)s for fandeck %(fandeck)s.',
                                    code=colour_code, fandeck=fandeck_name or '-'))
        category = self.env['product.category'].search([('name', '=ilike', category_name)], limit=2)
        if len(category) != 1:
            raise ValidationError(_('Base Category "%s" was not found uniquely.', category_name))
        uom = self.env['uom.uom'].search([('name', '=ilike', uom_name)], limit=2)
        if len(uom) != 1:
            raise ValidationError(_('Base UoM "%s" was not found uniquely.', uom_name))
        normalized_attribute = service.normalize_attribute(attribute_name)
        attribute = False
        if normalized_attribute:
            candidates = self.env['product.template.attribute.value'].search([])
            attribute = candidates.filtered(
                lambda item: service.normalize_attribute(item.name) == normalized_attribute
            )[:1]
            if not attribute:
                raise ValidationError(_('Base Attribute "%s" was not found.', attribute_name))
        colour = colour_candidates[:1]
        if len(colour_candidates) > 1:
            families = self.env['tinting.formula'].search([
                ('colour_code_id', 'in', colour_candidates.ids),
                ('base_category_id', '=', category.id),
                ('base_uom_id', '=', uom.id),
            ])
            families = families.filtered(
                lambda item: service.normalize_attribute(item.base_attribute_name) == normalized_attribute
            )
            family_colours = families.mapped('colour_code_id')
            if len(family_colours) == 1:
                colour = family_colours
                _logger.warning(
                    "FORMULA_IMPORT_DUPLICATE_COLOUR_MASTER_RESOLVED code=%s fandeck=%s "
                    "candidate_ids=%s selected_id=%s family_ids=%s",
                    colour_code, fandeck_name, colour_candidates.ids, colour.id, families.ids,
                )
            else:
                raise ValidationError(_(
                    'Colour code %(code)s has multiple master records in fandeck %(fandeck)s. '
                    'Clean the duplicate colour masters before importing a new formula.',
                    code=colour_code, fandeck=fandeck_name or '-',
                ))
        return {
            'colour': colour,
            'category': category,
            'uom': uom,
            'attribute': attribute,
            'attribute_name': normalized_attribute,
        }

    @api.model
    def create_draft_from_row(self, row):
        service = self.env['tinting.formula.recipe.service']
        resolved = self._resolve_identity(
            row.fandeck_name, row.colour_code, row.base_category_name,
            row.base_uom_name, row.base_attribute_name,
        )
        recipe = json.loads(row.recipe_json or '{}')
        family = service.find_family(
            resolved['colour'], resolved['category'], resolved['uom'], resolved['attribute_name'],
        )
        if not family:
            family = self.env['tinting.formula'].create({
                'colour_code_id': resolved['colour'].id,
                'base_category_id': resolved['category'].id,
                'base_uom_id': resolved['uom'].id,
                'base_attribute_id': resolved['attribute'].id if resolved['attribute'] else False,
                'state': 'draft',
                'notes': row.notes,
            })
        else:
            service.ensure_legacy_version(family)
        signature = service.signature(recipe)
        duplicate = family.version_ids.filtered(lambda item: item.recipe_signature == signature)[:1]
        if duplicate:
            row.write({
                'state': 'duplicate', 'planned_action': 'skip',
                'message': _('Recipe became a duplicate of %s before import.', duplicate.display_name),
                'formula_id': family.id,
            })
            _logger.warning(
                "FORMULA_IMPORT_RACE_DUPLICATE row_id=%s family_id=%s duplicate_version_id=%s",
                row.id, family.id, duplicate.id,
            )
            return duplicate
        previous = family.latest_version_id or family.default_version_id
        version = self.env['tinting.formula.version'].create_from_recipe(
            family=family,
            recipe=recipe,
            state='draft',
            origin='import',
            previous_version=previous,
            import_batch=row.batch_id,
            notes=row.notes,
            change_reason=row.change_reason or _('Spreadsheet import %(reference)s',
                                                reference=row.external_reference or row.row_number),
        )
        row.write({
            'state': 'imported', 'message': _('Draft %s created.', version.display_name),
            'formula_id': family.id, 'version_id': version.id,
        })
        _logger.info(
            "FORMULA_IMPORT_DRAFT_CREATED batch_id=%s row=%s family_id=%s version_id=%s",
            row.batch_id.id, row.row_number, family.id, version.id,
        )
        return version
